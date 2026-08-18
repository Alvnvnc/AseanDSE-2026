#!/usr/bin/env python3
"""Ekspor semua data proyek ke berkas Excel (.xlsx) berbahasa Inggris.

Menghasilkan dua hal di `ekspor-xlsx/`:

1. `ASEAN_DSE_hasil_analisis.xlsx` — satu buku kerja berisi seluruh tabel hasil
   analisis (`analisis/keluaran/*.csv`) plus lembar "Contents" dan "Key figures".
   Ini yang dipakai menulis naskah storyboard / bikin grafik cepat di Excel.
2. `sac/<nama>.xlsx` — satu berkas per himpunan data mentah (`data/siap-sac/*.csv`),
   karena SAP Analytics Cloud mengimpor satu himpunan data per berkas.

Storyboard dinilai dalam bahasa Inggris, jadi SELURUH isi .xlsx (nama lembar,
nama kolom, isi sel kategori, dan lembar angka kunci) diterjemahkan di sini lewat
`istilah_en.py`. CSV kerja tetap berbahasa Indonesia — notebook dan `siapkan.py`
tidak ikut berubah. Nama diri (provinsi, kab/kota, kota) tidak diterjemahkan.

Jalankan: .venv/bin/python analisis/ekspor_xlsx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import istilah_en as en  # noqa: E402

AKAR = Path(__file__).resolve().parent.parent
KELUARAN = AKAR / "analisis" / "keluaran"
SIAP_SAC = AKAR / "data" / "siap-sac"
TUJUAN = AKAR / "ekspor-xlsx"

WARNA_KEPALA = PatternFill("solid", fgColor="1F4E79")
FONT_KEPALA = Font(bold=True, color="FFFFFF")

# Urutan lembar = urutan halaman storyboard. Nama lembar & keterangan berbahasa
# Inggris karena langsung terbaca juri lewat storyboard.
# asal "csv:<nama>" dibaca dari analisis/keluaran/, "hitung" dihasilkan lembar_grafik()
# (definisi identik dengan analisis_storyboard.ipynb).
URUTAN: list[tuple[str, str, str, str]] = [
    ("csv:dengue_asean_nasional_tahunan", "Dengue national annual", "2",
     "Dengue cases by country and year (9 ASEAN countries); the 'complete' column flags the years that are safe to compare."),
    ("hitung", "Seasonality IDN-THA", "5",
     "Mean cases, rainfall and temperature per calendar month (complete years from 2010) — the seasonal curve."),
    ("csv:korelasi_jeda_nasional", "Lag correlation - national", "6",
     "Correlation of cases against temperature/rainfall/ONI at lags of 0-8 months, raw series (seasonality still in)."),
    ("hitung", "Monthly series IDN-THA", "6, 7, 10, 11",
     "The backbone: national monthly series from 2010 — cases, anomaly, lagged ONI/temperature/rainfall, P75 threshold, alarm status."),
    ("hitung", "ENSO phase - averages", "7",
     "The ace card: mean cases per month in Indonesia by the ENSO phase 4 months earlier (strong El Nino vs neutral)."),
    ("csv:korelasi_jeda_anomali", "Lag correlation - anomaly", "7",
     "The same correlations but on ANOMALIES (seasonal effect removed) — evidence the signal is more than seasonality."),
    ("hitung", "Province consistency IDN", "7",
     "Ratio of cases in El Nino months to other months per province — evidence the effect is not an artefact of national aggregation."),
    ("csv:kalender_risiko_idn", "Risk calendar IDN", "8, 10",
     "Share of cases per month per Indonesian province + Peak/Watch/Normal status (status_code 2/1/0)."),
    ("csv:kalender_risiko_tha", "Risk calendar THA", "8",
     "The same risk calendar for Thailand's 77 provinces (cross-country replication)."),
    ("csv:insidens_provinsi_2018_2020", "Province incidence 18-20", "9, 12",
     "Incidence per 100k population, 2018-2020 ONLY (official BPS denominator). Basis for geographic prioritisation."),
    ("hitung", "City flood vulnerability", "9",
     "The 25 ASEAN cities with the largest population inside the 100-year flood zone (2020 & 2030 projection)."),
    ("hitung", "WASH by country", "9",
     "Piped water, drinking water and sanitation access by country (latest year) — the SDG 6 layer."),
    ("csv:ambang_pemicu_evaluasi", "Trigger rules - evaluation", "11",
     "Alarm rules A-D compared on the full 2010-2024 data: coverage, sensitivity, precision, lift."),
    ("csv:ambang_pemicu_luar_sampel", "Trigger rules - out of sample", "12",
     "Blind test: rules built on 2010-2016, tested on 2017-2024 (Indonesia & Thailand)."),
    ("csv:ambang_pemicu_vs_kanal_endemis", "Rules vs endemic channel", "12",
     "The rules tested against the standard WHO/PAHO outbreak definition (endemic channel, 5-year mean+2SD)."),
    ("csv:benchmark_ews", "EWS benchmark", "12",
     "Our rule's performance set beside published dengue early warning systems."),
    ("csv:episode_elnino_tenggang", "El Nino episodes - lead time", "13",
     "Every El Nino episode since 2009: when the alarm fires, when cases cross P75, how many months of lead time that buys."),
    ("csv:dampak_moneter", "Monetary impact", "13",
     "Scenarios of 10/20/30% of peak-season cases averted -> value in US$ and Rp per year."),
    ("csv:inovasi_kekeringan_akses_air", "Drought x water access", "14",
     "Test of the drought x water-access innovation angle — EXPLORATORY (n=7) and DROPPED as a causal claim."),
    ("csv:banjir_vs_dengue_uji", "Flood vs dengue tests", "14",
     "Nine tests of the flood/disaster -> dengue link. The raw correlation (+0.66) is a population artefact: it vanishes per capita (-0.14). Reported as a NULL result."),
    ("csv:banjir_vs_dengue_panel", "Flood vs dengue panel", "14",
     "The province-year panel behind those tests (33 provinces x 2018-2020) — feeds the raw-vs-per-capita scatter pair."),
]

KETERANGAN_SAC: dict[str, str] = {
    "dengue_asean": "Raw ASEAN dengue cases from OpenDengue (all space & time resolutions).",
    "dengue_iklim_bulanan": "Main analysis table: monthly cases + temperature/rainfall/ONI at lags 0-4 months.",
    "iklim_bulanan_asean": "Monthly temperature and rainfall by country, 1950-present.",
    "enso_oni_bulanan": "Monthly NOAA/CPC ONI index + ENSO phase (official 0.5 threshold).",
    "penduduk_provinsi_indonesia": "Indonesian province population (BPS WebAPI) — incidence denominator 2018-2020.",
    "wash_rumahtangga_asean": "Household drinking water & sanitation access (JMP WHO/UNICEF).",
    "wash_fasyankes_asean": "Water & sanitation access in health-care facilities (JMP WHO/UNICEF).",
    "bencana_kabkota_indonesia": "Disaster events per Indonesian district/city (BNPB).",
    "sampah_kabkota_indonesia": "Waste generation & management per district/city (SIPSN KLHK).",
    "kota_asean_paparan_banjir": "ASEAN city population exposed to flooding (GHSL).",
    "kota_asean_kualitas_udara": "PM2.5/PM10/NO2 concentrations in ASEAN cities (WHO Ambient Air Quality).",
}

POLA_RIBUAN = re.compile(r"^-?\d{1,3}(,\d{3})+$")


def rapikan(df: pd.DataFrame, nama: str) -> pd.DataFrame:
    """Beri nama kolom tanpa judul dan ubah angka ber-pemisah-ribuan jadi numerik."""
    ganti = {"ambang_pemicu_luar_sampel": "periode", "ambang_pemicu_vs_kanal_endemis": "aturan"}
    df = df.rename(columns={c: ganti.get(nama, "keterangan")
                            for c in df.columns if str(c).startswith("Unnamed:")})
    for kol in df.columns:
        if df[kol].dtype == "object" or pd.api.types.is_string_dtype(df[kol]):
            teks = df[kol].dropna().astype(str)
            if len(teks) and teks.map(lambda v: bool(POLA_RIBUAN.match(v))).all():
                df[kol] = pd.to_numeric(df[kol].astype(str).str.replace(",", "", regex=False))
    return df


def tulis(df: pd.DataFrame, penulis: pd.ExcelWriter, lembar: str, bekukan: bool = True) -> None:
    """Tulis satu lembar; penerjemahan ke bahasa Inggris terjadi di sini."""
    df = en.terjemahkan(df)
    df.to_excel(penulis, sheet_name=lembar, index=False)
    ws = penulis.book[lembar]
    for sel in ws[1]:
        sel.fill = WARNA_KEPALA
        sel.font = FONT_KEPALA
        sel.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    if bekukan:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for i, kol in enumerate(df.columns, start=1):
        contoh = df[kol].dropna().astype(str).head(200).str.len().max()
        lebar = max(len(str(kol)) + 4, int(contoh if pd.notna(contoh) else 8) + 2)
        ws.column_dimensions[get_column_letter(i)].width = min(max(lebar, 10), 42)


def baca_angka_kunci() -> pd.DataFrame:
    berkas = KELUARAN / "angka_kunci_storyboard.md"
    bagian, baris = "", []
    for teks in berkas.read_text(encoding="utf-8").splitlines():
        teks = teks.rstrip()
        if teks.startswith("## "):
            bagian = teks[3:].strip()
        elif teks.startswith("- "):
            baris.append({"bagian": bagian, "poin": teks[2:].strip()})
        elif teks.startswith("  ") and baris:  # lanjutan butir sebelumnya
            baris[-1]["poin"] += " " + teks.strip()
    df = pd.DataFrame(baris)
    df["poin"] = df["poin"].str.replace("**", "", regex=False).str.replace("`", "", regex=False)
    return df.map(en.prosa)


NAMA_BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
FASE_URUT = ["La Nina kuat", "La Nina", "Netral", "El Nino", "El Nino kuat"]


def seri_nasional(iso: str, d, ik, oni):
    """Deret bulanan nasional — definisi identik dengan notebook (sel 3)."""
    kasus = d[d["iso3"] == iso].groupby("periode")["kasus"].sum().rename("kasus")
    iklim = ik[ik["iso3"] == iso].set_index("periode")[["suhu_c", "hujan_mm"]]
    ens = oni.set_index("periode")[["oni", "fase"]]
    s = pd.concat([kasus, iklim, ens], axis=1)
    s.index = pd.PeriodIndex(s.index, freq="M")
    s = s.sort_index()
    ada = s[s["kasus"].notna()].index
    s = s.reindex(pd.period_range(ada.min(), ada.max(), freq="M"))
    s["tahun"], s["bulan"] = s.index.year, s.index.month
    return s


def tahun_lengkap(s, mulai: int = 2010) -> list[int]:
    n = s[s["kasus"].notna() & (s["tahun"] >= mulai)].groupby("tahun")["bulan"].nunique()
    return sorted(n[n == 12].index)


def lembar_grafik() -> dict[str, pd.DataFrame]:
    """Hitung ulang lembar siap-grafik dari data sumber di data/siap-sac/."""
    d = pd.read_csv(SIAP_SAC / "dengue_iklim_bulanan.csv", low_memory=False)
    ik = pd.read_csv(SIAP_SAC / "iklim_bulanan_asean.csv")
    oni = pd.read_csv(SIAP_SAC / "enso_oni_bulanan.csv")
    hasil: dict[str, pd.DataFrame] = {}

    # --- bulan puncak Indonesia (4 bulan dengan rata-rata kasus tertinggi) ---
    s_idn = seri_nasional("IDN", d, ik, oni)
    th_idn = tahun_lengkap(s_idn)
    puncak_idn = sorted(s_idn[s_idn["tahun"].isin(th_idn)].groupby("bulan")["kasus"].mean()
                        .nlargest(4).index)

    # --- 1. deret bulanan nasional IDN & THA ---
    deret = []
    for iso, nama in [("IDN", "Indonesia"), ("THA", "Thailand")]:
        s = seri_nasional(iso, d, ik, oni)
        s["suhu_jeda2"] = s["suhu_c"].shift(2)
        s["hujan_jeda1"] = s["hujan_mm"].shift(1)
        s["oni_jeda3"], s["oni_jeda4"] = s["oni"].shift(3), s["oni"].shift(4)
        s["fase_jeda4"] = s["fase"].shift(4)
        s = s[s.index >= "2010-01"].dropna(subset=["kasus"]).copy()
        p75 = s["kasus"].quantile(0.75)
        rer = s.groupby("bulan")["kasus"].transform("mean")
        sd = s.groupby("bulan")["kasus"].transform("std")
        s["kasus_anomali_z"] = (s["kasus"] - rer) / sd
        s["negara"] = nama
        s["periode"] = s.index.astype(str)
        s["tanggal"] = s.index.to_timestamp()
        s["nama_bulan"] = [NAMA_BULAN[b - 1] for b in s["bulan"]]
        s["ambang_P75"] = round(p75)
        s["bulan_puncak_kasus"] = (s["kasus"] >= p75).map({True: "Ya", False: "Tidak"})
        if iso == "IDN":
            alarm = s["bulan"].isin(puncak_idn) & (s["oni_jeda4"] >= 0.5)
            s["alarm_aturan_D"] = alarm.map({True: "Ya", False: "Tidak"})
            s["alarm_El_Nino"] = (s["oni_jeda4"] >= 0.5).map({True: "Ya", False: "Tidak"})
        else:  # Thailand memakai anomali suhu jeda-2, bukan ONI
            s["alarm_aturan_D"] = "n/a"
            s["alarm_El_Nino"] = (s["oni_jeda4"] >= 0.5).map({True: "Ya", False: "Tidak"})
        deret.append(s)
    kol = ["negara", "periode", "tanggal", "tahun", "bulan", "nama_bulan", "kasus",
           "kasus_anomali_z", "ambang_P75", "bulan_puncak_kasus", "alarm_aturan_D",
           "alarm_El_Nino", "fase_jeda4", "oni", "oni_jeda3", "oni_jeda4",
           "suhu_c", "suhu_jeda2", "hujan_mm", "hujan_jeda1"]
    dr = pd.concat(deret)[kol].reset_index(drop=True)
    num = dr.select_dtypes("number").columns
    dr[num] = dr[num].round(3)
    hasil["Monthly series IDN-THA"] = dr

    # --- 2. profil musiman (tahun lengkap 2010+) ---
    baris = []
    for iso, nama in [("IDN", "Indonesia"), ("THA", "Thailand")]:
        s = seri_nasional(iso, d, ik, oni)
        v = s[s["tahun"].isin(tahun_lengkap(s))]
        g = v.groupby("bulan").agg(kasus_rerata=("kasus", "mean"),
                                   hujan_rerata_mm=("hujan_mm", "mean"),
                                   suhu_rerata_c=("suhu_c", "mean"))
        th = tahun_lengkap(s)
        for b, r in g.iterrows():
            baris.append({"negara": nama, "bulan": b, "nama_bulan": NAMA_BULAN[b - 1],
                          "kasus_rerata": round(r["kasus_rerata"]),
                          "hujan_rerata_mm": round(r["hujan_rerata_mm"], 1),
                          "suhu_rerata_c": round(r["suhu_rerata_c"], 2),
                          "jendela_tahun_lengkap": f"{th[0]}-{th[-1]} (n={len(th)})"})
    hasil["Seasonality IDN-THA"] = pd.DataFrame(baris)

    # --- 3. rata-rata kasus menurut fase ENSO jeda-4 (kartu as halaman 7) ---
    s = seri_nasional("IDN", d, ik, oni)
    s["fase_j4"] = s["fase"].shift(4)
    v = s[s.index >= "2010-01"].dropna(subset=["kasus", "fase_j4"])
    tab = v.groupby("fase_j4")["kasus"].agg(["mean", "count"]).reindex(FASE_URUT)
    netral = tab.loc["Netral", "mean"]
    hasil["ENSO phase - averages"] = pd.DataFrame({
        "fase_enso_jeda4": FASE_URUT,
        "urutan": range(1, 6),
        "rerata_kasus_per_bulan": tab["mean"].round().astype(int).values,
        "n_bulan": tab["count"].astype(int).values,
        "rasio_thd_netral": (tab["mean"] / netral).round(2).values,
        "cukup_data": [f"Tidak (n={n:.0f})" if n < 10 else "Ya" for n in tab["count"]],
    })

    # --- 4. konsistensi antar provinsi Indonesia ---
    p = d[(d["iso3"] == "IDN") & (d["tahun"] >= 2010)].dropna(subset=["kasus", "fase_enso_jeda4"])
    baris = []
    for prov, g in p.groupby("provinsi"):
        nino = g["fase_enso_jeda4"].isin(["El Nino", "El Nino kuat"])
        if nino.sum() >= 12 and (~nino).sum() >= 24 and g.loc[~nino, "kasus"].mean() > 0:
            baris.append({"provinsi": str(prov).title(),
                          "rasio_nino_vs_lain": round(g.loc[nino, "kasus"].mean() /
                                                      g.loc[~nino, "kasus"].mean(), 2),
                          "n_bulan_nino": int(nino.sum()),
                          "arah": "Naik saat El Nino" if g.loc[nino, "kasus"].mean() >
                                  g.loc[~nino, "kasus"].mean() else "Tidak naik"})
    hasil["Province consistency IDN"] = (pd.DataFrame(baris)
                                        .sort_values("rasio_nino_vs_lain", ascending=False)
                                        .reset_index(drop=True))

    # --- 5. paparan banjir kota (halaman 9) ---
    b = pd.read_csv(SIAP_SAC / "kota_asean_paparan_banjir.csv")
    piv = (b[b["indikator"] == "Penduduk terpapar banjir 100 tahunan"]
           .pivot_table(index=["negara", "kota", "penduduk_kota_2025"],
                        columns="tahun", values="nilai").reset_index())
    piv = piv.rename(columns={2020: "terpapar_2020", 2030: "terpapar_2030"})
    # pakai indikator persen bawaan GHSL (penyebut = penduduk tahun itu), bukan hitungan sendiri
    pers = (b[(b["indikator"] == "Persen penduduk terpapar banjir 100 tahunan") & (b["tahun"] == 2020)]
            .set_index(["negara", "kota"])["nilai"].rename("persen_2020"))
    piv = piv.join(pers, on=["negara", "kota"])
    piv["pertumbuhan_2020_2030_%"] = ((piv["terpapar_2030"] / piv["terpapar_2020"] - 1) * 100)
    piv = piv[piv["penduduk_kota_2025"] >= 500_000].nlargest(25, "terpapar_2020")
    piv[["penduduk_kota_2025", "terpapar_2020", "terpapar_2030"]] = \
        piv[["penduduk_kota_2025", "terpapar_2020", "terpapar_2030"]].round(0)
    piv[["persen_2020", "pertumbuhan_2020_2030_%"]] = piv[["persen_2020", "pertumbuhan_2020_2030_%"]].round(1)
    hasil["City flood vulnerability"] = piv[["negara", "kota", "penduduk_kota_2025",
                                            "terpapar_2020", "persen_2020", "terpapar_2030",
                                            "pertumbuhan_2020_2030_%"]].reset_index(drop=True)

    # --- 6. WASH per negara, tahun terbaru (halaman 9) ---
    w = pd.read_csv(SIAP_SAC / "wash_rumahtangga_asean.csv")
    w = w[(w["wilayah"] == "Total") & w["persen"].notna()]
    pilih = [("Air minum", "Perpipaan", "air_perpipaan_%"),
             ("Air minum", "Minimal dasar", "air_minimal_dasar_%"),
             ("Air minum", "Dikelola aman", "air_dikelola_aman_%"),
             ("Sanitasi", "Minimal dasar", "sanitasi_minimal_dasar_%")]
    tab, tahun_maks = None, {}
    for layanan, ind, nama_kol in pilih:
        sub = w[(w["layanan"] == layanan) & (w["indikator"] == ind)]
        sub = sub.loc[sub.groupby("negara")["tahun"].idxmax()][["negara", "tahun", "persen"]]
        tahun_maks[nama_kol] = int(sub["tahun"].max())
        sub = sub.rename(columns={"persen": nama_kol}).drop(columns="tahun")
        tab = sub if tab is None else tab.merge(sub, on="negara", how="outer")
    tab.insert(1, "tahun_data", max(tahun_maks.values()))
    tab = tab[["negara", "tahun_data"] + [k for _, _, k in pilih]]
    hasil["WASH by country"] = tab.sort_values("air_perpipaan_%").round(1).reset_index(drop=True)
    return hasil


def buku_analisis() -> Path:
    keluar = TUJUAN / "ASEAN_DSE_hasil_analisis.xlsx"
    isi = []
    with pd.ExcelWriter(keluar, engine="openpyxl") as penulis:
        # tempat penampung; daftar isi ditulis ulang di akhir agar berada di depan
        pd.DataFrame({"a": [0]}).to_excel(penulis, sheet_name="Contents", index=False)
        tulis(baca_angka_kunci(), penulis, "Key figures")
        penulis.book["Key figures"].column_dimensions["B"].width = 110
        for sel in penulis.book["Key figures"]["B"]:
            sel.alignment = Alignment(wrap_text=True, vertical="top")

        grafik = lembar_grafik()
        for asal, lembar, halaman, ket in URUTAN:
            if asal == "hitung":
                df, sumber = grafik[lembar], "recomputed from data/siap-sac/"
            else:
                nama = asal.split(":", 1)[1]
                df, sumber = rapikan(pd.read_csv(KELUARAN / f"{nama}.csv"), nama), f"{nama}.csv"
            tulis(df, penulis, lembar)
            isi.append({"sheet": lembar, "page": halaman, "contents": ket,
                        "rows": len(df), "columns": df.shape[1], "source file": sumber})

        del penulis.book["Contents"]
        df_isi = pd.DataFrame(
            [{"sheet": "Key figures", "page": "all",
              "contents": "Every number quoted in the storyboard, grouped by storyboard page.",
              "rows": "-", "columns": "-", "source file": "angka_kunci_storyboard.md"}] + isi)
        tulis(df_isi, penulis, "Contents", bekukan=False)
        penulis.book.move_sheet("Contents", offset=-(len(penulis.book.sheetnames) - 1))
        penulis.book["Contents"].column_dimensions["C"].width = 95
        for sel in penulis.book["Contents"]["C"]:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    return keluar


def buku_sac() -> list[tuple[str, int, int, float]]:
    folder = TUJUAN / "sac"
    folder.mkdir(parents=True, exist_ok=True)
    hasil = []
    for csv in sorted(SIAP_SAC.glob("*.csv")):
        df = pd.read_csv(csv, low_memory=False)
        keluar = folder / f"{csv.stem}.xlsx"
        with pd.ExcelWriter(keluar, engine="openpyxl") as penulis:
            tulis(df, penulis, csv.stem[:31])
        hasil.append((csv.stem, len(df), df.shape[1], keluar.stat().st_size / 1e6))
        print(f"  {keluar.name:<45} {len(df):>7,} baris  {keluar.stat().st_size/1e6:>5.1f} MB")
        print(f"  {'':<45} {KETERANGAN_SAC.get(csv.stem, '')}")
    return hasil


if __name__ == "__main__":
    TUJUAN.mkdir(parents=True, exist_ok=True)
    print("Buku kerja hasil analisis (isi berbahasa Inggris):")
    berkas = buku_analisis()
    print(f"  {berkas.relative_to(AKAR)}  ({berkas.stat().st_size/1e6:.1f} MB)")
    print("\nHimpunan data untuk SAC (satu berkas per himpunan data):")
    buku_sac()
    print("\nSelesai.")
